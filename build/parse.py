import docx, json, re
from docx.document import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
def blocks(d):
    for el in d.element.body.iterchildren():
        if el.tag.endswith('}p'): yield Paragraph(el,d)
        elif el.tag.endswith('}tbl'): yield Table(el,d)
def clean(s): return re.sub(r'\s+',' ',s).strip()
def bullets(cell):
    items=[]
    for p in cell.paragraphs:
        t=p.text
        if not clean(t): continue
        parts=[clean(x) for x in re.split(r'—[\s\u00a0]{2,}',t) if clean(x)]
        items+=parts
    return items
def rows(tbl):
    out={}
    for r in tbl.rows:
        cs=r.cells
        if len(cs)>=2: out[clean(cs[0].text)]=cs[1]
    return out

U='../source/'
# Program
d=docx.Document(U+'Team_Leader_Raising_Program_2.docx')
comps=[];cur=None;week=None
for b in blocks(d):
    if isinstance(b,Paragraph):
        t=clean(b.text)
        m=re.match(r'^(\d{1,2})\s+—\s+(.+)$',t)
        if m: cur={'id':int(m[1]),'title':m[2],'week':week}; comps.append(cur)
    else:
        txt=clean(b.rows[0].cells[0].text)
        m=re.match(r'^(WEEKS? [\dTO ]+|CHECKPOINT|FINAL)',txt)
        if m and len(b.rows)==1 and len(b.columns)==1:
            week=txt; cur=None; continue
        if cur is None: continue
        r=rows(b)
        if 'Purpose' in r and 'title' in cur and 'purpose' not in cur:
            cur['purpose']=clean(r['Purpose'].text)
            if 'Uses' in r: cur['uses']=clean(r['Uses'].text)
            if 'What good looks like' in r: cur['good']=bullets(r['What good looks like'])
            if 'Common failures' in r: cur['fail']=bullets(r['Common failures'])
            cur['task']=clean(r['Practice task'].text)
            cur['signoff']=clean(r['Sign-off'].text)
json.dump(comps,open('comps.json','w'),indent=1,ensure_ascii=False)
for c in comps: print(c['id'],c['title'],'|',c['week'][:40],'|',len(c.get('good',[])),len(c.get('fail',[])),c.get('uses'))

# Scenarios
d=docx.Document(U+'Phase_1_Scenario_Test_Candidate.docx')
sc=[];cur=None
for b in blocks(d):
    if isinstance(b,Paragraph):
        m=re.match(r'^(\d{1,2})\s+—\s+(.+)$',clean(b.text))
        if m: cur={'id':int(m[1]),'title':m[2]}; sc.append(cur)
    elif cur and 'situation' not in cur:
        r=rows(b)
        if 'Situation' in r:
            cur['situation']=clean(r['Situation'].text)
            q=clean(r['Questions'].text)
            cur['questions']=[clean(x) for x in re.split(r'(?=\b[a-d]\)\s)',q) if clean(x)]
json.dump(sc,open('scen.json','w'),indent=1,ensure_ascii=False)
for s in sc: print(s['id'],s['title'],s['questions'])

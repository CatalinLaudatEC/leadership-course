import docx, json, re, sys
from docx.table import Table
from docx.text.paragraph import Paragraph
U='../source/'
def blocks(parent, body):
    for el in body.iterchildren():
        if el.tag.endswith('}p'): yield Paragraph(el,parent)
        elif el.tag.endswith('}tbl'): yield Table(el,parent)
def runs(p):
    segs=[]
    for r in p.runs:
        t=r.text
        if not t: continue
        b=bool(r.bold); i=bool(r.italic)
        if segs and segs[-1]['b']==b and segs[-1]['i']==i: segs[-1]['t']+=t
        else: segs.append({'t':t,'b':b,'i':i})
    # images
    img = bool(p._p.xpath('.//pic:pic'))
    return segs, img
def para(p):
    segs,img=runs(p)
    txt=''.join(s['t'] for s in segs)
    st=p.style.name if p.style is not None else ''
    lvl=None
    try: lvl=p._p.pPr.numPr is not None
    except: pass
    return {'k':'p','style':st,'segs':segs,'text':txt,'img':img,'list':bool(lvl)}
def table(t):
    rows=[]
    for r in t.rows:
        seen=set(); cells=[]
        for c in r.cells:
            if id(c._tc) in seen: continue
            seen.add(id(c._tc))
            items=[]
            for b in blocks(c, c._tc):
                items.append(para(b) if isinstance(b,Paragraph) else table(b))
            cells.append({'span':len([x for x in r.cells if x._tc is c._tc]),'items':items})
        rows.append(cells)
    return {'k':'tbl','rows':rows}
out={}
for f in ['Team_Leader_Raising_Program_2','Team_Communication_Workflow','Project_Technical_Standards','Phase_1_Scenario_Test_Candidate','Phase_1_Scenario_Test_Mentor']:
    d=docx.Document(U+f+'.docx')
    bl=[]
    for b in blocks(d,d.element.body):
        x=para(b) if isinstance(b,Paragraph) else table(b)
        if x['k']=='p' and not x['text'].strip() and not x['img']: continue
        bl.append(x)
    # header/footer
    hdr=[]
    for s in d.sections:
        for p in s.header.paragraphs:
            if p.text.strip(): hdr.append(p.text)
        for p in s.footer.paragraphs:
            if p.text.strip(): hdr.append('FOOTER:'+p.text)
    out[f]={'blocks':bl,'hdr':hdr}
json.dump(out,open('docs_raw.json','w'),ensure_ascii=False)
# summary
def sm(x,ind=0):
    pre=' '*ind
    if x['k']=='p':
        print(pre+'P[%s]%s%s %s'%(x['style'],' IMG' if x['img'] else '',' LI' if x['list'] else '', ''.join(('**'+s['t']+'**') if s['b'] else s['t'] for s in x['segs'])[:110]))
    else:
        print(pre+'TBL %dx%s'%(len(x['rows']),[len(r) for r in x['rows']][:3]))
        for r in x['rows'][:2]:
            for c in r:
                for it in c['items'][:3]: sm(it,ind+4)
for f,v in out.items():
    print('#######',f,v['hdr'])
    for b in v['blocks']: sm(b)
